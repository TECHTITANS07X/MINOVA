from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any

import structlog

logger = structlog.get_logger()


@dataclass
class MetricRow:
    name: str
    value: Decimal
    unit: str
    target: Decimal | None = None
    variance: Decimal | None = None
    variance_pct: Decimal | None = None


@dataclass
class ReportData:
    mine_name: str
    period: str
    start_date: date
    end_date: date
    metrics: list[MetricRow]
    summary_text: str = ""


def generate_excel(report: ReportData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    header_font = Font(bold=True, size=14)
    sub_font = Font(bold=True, size=11)
    th_fill = PatternFill("solid", fgColor="1F4E79")
    th_font = Font(bold=True, color="FFFFFF", size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    ws.merge_cells("A1:F1")
    ws["A1"] = f"MINOVA — {report.mine_name}"
    ws["A1"].font = header_font
    ws.merge_cells("A2:F2")
    ws["A2"] = f"{report.period} | {report.start_date} to {report.end_date}"
    ws["A2"].font = sub_font

    headers = ["Metric", "Value", "Unit", "Target", "Variance", "Var %"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = th_font
        cell.fill = th_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for i, m in enumerate(report.metrics, 5):
        ws.cell(row=i, column=1, value=m.name).border = thin_border
        ws.cell(row=i, column=2, value=float(m.value)).border = thin_border
        ws.cell(row=i, column=3, value=m.unit).border = thin_border
        ws.cell(row=i, column=4, value=float(m.target) if m.target else "").border = thin_border
        ws.cell(row=i, column=5, value=float(m.variance) if m.variance else "").border = thin_border
        ws.cell(row=i, column=6, value=f"{m.variance_pct}%" if m.variance_pct else "").border = thin_border

    if report.summary_text:
        row = len(report.metrics) + 6
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1, value=report.summary_text)

    for col in range(1, 7):
        ws.column_dimensions[chr(64 + col)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf(report: ReportData) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"MINOVA — {report.mine_name}", styles["Title"]))
    elements.append(Paragraph(f"{report.period} | {report.start_date} to {report.end_date}", styles["Heading2"]))
    elements.append(Spacer(1, 12))

    data = [["Metric", "Value", "Unit", "Target", "Variance", "Var %"]]
    for m in report.metrics:
        data.append([
            m.name,
            f"{m.value:,.2f}",
            m.unit,
            f"{m.target:,.2f}" if m.target else "-",
            f"{m.variance:,.2f}" if m.variance else "-",
            f"{m.variance_pct}%" if m.variance_pct else "-",
        ])

    table = Table(data, colWidths=[80, 60, 40, 60, 60, 50])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
    ]))
    elements.append(table)

    if report.summary_text:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph(report.summary_text, styles["Normal"]))

    doc.build(elements)
    return buf.getvalue()


def generate_docx(report: ReportData) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10)

    title = doc.add_heading(f"MINOVA — {report.mine_name}", level=1)
    doc.add_paragraph(f"{report.period} | {report.start_date} to {report.end_date}")
    doc.add_paragraph("")

    table = doc.add_table(rows=1 + len(report.metrics), cols=6)
    table.style = "Light Grid Accent 1"
    headers = ["Metric", "Value", "Unit", "Target", "Variance", "Var %"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    for row_idx, m in enumerate(report.metrics, 1):
        table.rows[row_idx].cells[0].text = m.name
        table.rows[row_idx].cells[1].text = f"{m.value:,.2f}"
        table.rows[row_idx].cells[2].text = m.unit
        table.rows[row_idx].cells[3].text = f"{m.target:,.2f}" if m.target else "-"
        table.rows[row_idx].cells[4].text = f"{m.variance:,.2f}" if m.variance else "-"
        table.rows[row_idx].cells[5].text = f"{m.variance_pct}%" if m.variance_pct else "-"

    if report.summary_text:
        doc.add_paragraph("")
        doc.add_paragraph(report.summary_text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
